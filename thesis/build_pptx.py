# -*- coding: utf-8 -*-
"""Build the bachelor-defense presentation (.pptx), ~12 slides, ~10 min.

Run from repo root:  python3 thesis/build_pptx.py
Output: thesis/Презентація_захист_Макаренко.pptx
Figures come from runs/figures/ and thesis/figs/.
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUNF = os.path.join(ROOT, "runs", "figures")
TFIG = os.path.join(ROOT, "thesis", "figs")
OUT = os.path.join(ROOT, "thesis", "Презентація_захист_Макаренко.pptx")

NAVY = RGBColor(0x19, 0x40, 0x7A)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xB0, 0x30, 0x30)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def _set(run, size, bold=False, color=DARK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def title_bar(slide, text, num=None):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(1.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.4); tf.margin_top = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    _set(r, 26, bold=True, color=WHITE)
    if num is not None:
        nb = slide.shapes.add_textbox(SW - Inches(1.0), SH - Inches(0.5),
                                      Inches(0.8), Inches(0.4))
        np = nb.text_frame.paragraphs[0]; np.alignment = PP_ALIGN.RIGHT
        nr = np.add_run(); nr.text = str(num); _set(nr, 12, color=GREY)


def body_box(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    return tf


def bullet(tf, text, size=18, level=0, bold=False, color=DARK, first=False,
           space=6, dash=True):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.space_after = Pt(space)
    prefix = ("" if level == 0 and not dash else ("• " if level == 0 else "– "))
    r = p.add_run(); r.text = prefix + text
    _set(r, size, bold=bold, color=color)
    return p


def add_image_fit(slide, path, left, top, max_w, max_h):
    from PIL import Image
    try:
        iw, ih = Image.open(path).size
    except Exception:
        iw, ih = (4, 3)
    ar = iw / ih
    w, h = max_w, int(max_w / ar)
    if h > max_h:
        h = max_h; w = int(max_h * ar)
    left2 = left + (max_w - w) // 2
    top2 = top + (max_h - h) // 2
    slide.shapes.add_picture(path, left2, top2, width=w, height=h)


# ============================================================ SLIDE 1 — title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
stripe = s.shapes.add_shape(1, 0, Inches(4.5), SW, Inches(0.06))
stripe.fill.solid(); stripe.fill.fore_color.rgb = RGBColor(0xF0, 0xC4, 0x19)
stripe.line.fill.background()
tf = body_box(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Київський національний університет імені Тараса Шевченка"
_set(r, 16, color=WHITE)
for t in ["Факультет комп’ютерних наук та кібернетики",
          "Кафедра теоретичної кібернетики"]:
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; _set(r, 14, color=RGBColor(0xD5, 0xDE, 0xEC))

tf = body_box(s, Inches(0.8), Inches(1.9), Inches(11.7), Inches(2.2))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = ("Уніфікована система обробки та класифікації ЕКГ-сигналів різних "
          "форматів з використанням методів машинного навчання")
_set(r, 30, bold=True, color=WHITE)
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(10)
r = p.add_run(); r.text = "Кваліфікаційна робота на здобуття ступеня бакалавра"
_set(r, 17, italic=True, color=RGBColor(0xF0, 0xC4, 0x19))

tf = body_box(s, Inches(1.2), Inches(4.8), Inches(10.9), Inches(2.4))
rows = [
    ("Виконав: ", "студент 4-го курсу Макаренко Вадим Анатолійович, група ТК-41"),
    ("Науковий керівник: ", "д.ф.-м.н., професор, член-кореспондент НАН України "
     "Крак Юрій Васильович"),
]
first = True
for a, b in rows:
    p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
    p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(8)
    r = p.add_run(); r.text = a; _set(r, 16, bold=True, color=WHITE)
    r = p.add_run(); r.text = b; _set(r, 16, color=WHITE)
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(14)
r = p.add_run(); r.text = "Київ — 2026"; _set(r, 16, bold=True, color=RGBColor(0xF0, 0xC4, 0x19))


# ============================================ SLIDE 2 — coursework -> diploma
s = prs.slides.add_slide(BLANK)
title_bar(s, "Від курсової роботи до диплома", 2)
tf = body_box(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.8))
bullet(tf, "Було (курсова робота):", size=20, bold=True, color=NAVY, first=True)
bullet(tf, "уніфікований парсер форматів, ансамблевий детектор R-піків, "
       "морфологічний аналіз, аналіз HRV", size=17, level=1)
bullet(tf, "класифікація аритмій — лише на основі правил; машинне навчання "
       "не використовувалося", size=17, level=1)
bullet(tf, "Обмеження:", size=20, bold=True, color=ACCENT)
bullet(tf, "правила погано охоплюють нетипові форми; точність на складних "
       "випадках обмежена; немає кількісної оцінки за стандартним протоколом",
       size=17, level=1)
bullet(tf, "Стало (диплом) — що нове:", size=20, bold=True, color=NAVY)
bullet(tf, "повноцінний конвеєр машинного навчання на базі MIT-BIH "
       "(міжпацієнтний протокол)", size=17, level=1)
bullet(tf, "дві наукові новизни: двопотоковий гібридний класифікатор і "
       "бенчмарк багатоформатної стійкості", size=17, level=1)
bullet(tf, "REST-API та інтерфейс командного рядка; модульний пакет ekg",
       size=17, level=1)


# ============================================ SLIDE 3 — problem & goal
s = prs.slides.add_slide(BLANK)
title_bar(s, "Проблема та мета", 3)
tf = body_box(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.8))
bullet(tf, "Проблема", size=20, bold=True, color=NAVY, first=True)
bullet(tf, "серцево-судинні захворювання — провідна причина смертності; "
       "холтерівські записи надто великі для ручного перегляду", size=17, level=1)
bullet(tf, "ЕКГ зберігається в десятках несумісних форматів (DICOM, WFDB, "
       "EDF, HL7 aECG, SCP-ECG, CSV); модель має приймати їх усі", size=17, level=1)
bullet(tf, "Мета", size=20, bold=True, color=NAVY)
bullet(tf, "єдина система, що приймає ЕКГ різних форматів, зводить їх до "
       "спільного подання й класифікує кожне скорочення методами машинного "
       "навчання", size=17, level=1)
bullet(tf, "кількісно довести, що якість класифікації зберігається після "
       "перетворення між форматами", size=17, level=1)


# ============================================ SLIDE 4 — data & protocol
s = prs.slides.add_slide(BLANK)
title_bar(s, "Дані, класи та протокол оцінювання", 4)
tf = body_box(s, Inches(0.5), Inches(1.25), Inches(7.0), Inches(5.9))
bullet(tf, "MIT-BIH Arrhythmia Database: 48 записів, 360 Гц, 2 відведення",
       size=17, first=True)
bullet(tf, "5 класів за стандартом AAMI:", size=17)
bullet(tf, "N — норма, S — надшлуночкові, V — шлуночкові, F — зливні, "
       "Q — невизначені", size=15, level=1)
bullet(tf, "Чесний міжпацієнтний поділ DS1/DS2 (de Chazal):", size=17)
bullet(tf, "пацієнти навчання й тесту не перетинаються → реальна "
       "узагальнюваність, а не запам’ятовування ритму", size=15, level=1)
bullet(tf, "Сильна незбалансованість (≈90 % — клас N):", size=17)
bullet(tf, "головна метрика — macro-F1, а не accuracy", size=15, level=1, color=ACCENT)
add_image_fit(s, os.path.join(RUNF, "example_beats.png"),
              Inches(7.6), Inches(1.4), Inches(5.4), Inches(5.4))


# ============================================ SLIDE 5 — architecture
s = prs.slides.add_slide(BLANK)
title_bar(s, "Архітектура уніфікованої системи", 5)
add_image_fit(s, os.path.join(TFIG, "architecture.png"),
              Inches(0.4), Inches(1.2), Inches(8.6), Inches(6.0))
tf = body_box(s, Inches(9.1), Inches(1.5), Inches(4.0), Inches(5.4))
bullet(tf, "Модульний пакет ekg", size=17, bold=True, color=NAVY, first=True)
bullet(tf, "спільне проміжне подання у форматі XML", size=15, level=1)
bullet(tf, "додати формат = один перетворювач", size=15, level=1)
bullet(tf, "Повторно використано перевірені компоненти курсової роботи",
       size=15, bold=True, color=NAVY)
bullet(tf, "фільтрація, детектор R-піків, морфологія, HRV", size=15, level=1)
bullet(tf, "Доступ: REST-API і CLI", size=15, bold=True, color=NAVY)


# ============================================ SLIDE 6 — features & models
s = prs.slides.add_slide(BLANK)
title_bar(s, "Ознаки та моделі", 6)
tf = body_box(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.8))
bullet(tf, "Сегментація: вікно 260 відліків навколо R-піка, z-нормалізація",
       size=17, first=True)
bullet(tf, "Ознаки одного удару (35, для двох відведень — 70):", size=17)
bullet(tf, "інтервали RR, морфологія QRS, статистика, вейвлет db4", size=15, level=1)
bullet(tf, "Моделі:", size=17)
bullet(tf, "класичні — XGBoost, випадковий ліс, SVM-RBF (над ознаками)",
       size=15, level=1)
bullet(tf, "глибокі — 1D-CNN, CNN-BiLSTM (над сирим сигналом)", size=15, level=1)
bullet(tf, "гібрид — пізнє злиття двох потоків (новизна 1)", size=15, level=1, color=ACCENT)
bullet(tf, "Базова лінія — класифікація на основі правил (з курсової роботи)",
       size=17)


# ============================================ SLIDE 7 — novelty 1
s = prs.slides.add_slide(BLANK)
title_bar(s, "Новизна 1 — двопотоковий гібридний класифікатор", 7)
add_image_fit(s, os.path.join(TFIG, "hybrid.png"),
              Inches(0.4), Inches(1.25), Inches(8.4), Inches(3.6))
tf = body_box(s, Inches(0.5), Inches(5.0), Inches(12.3), Inches(2.3))
bullet(tf, "Ідея: пізнє злиття (stacking) — XGBoost над ручними ознаками + "
       "1D-CNN над сирим сигналом; мета-класифікатор (логістична регресія) на "
       "імовірностях обох потоків", size=16, first=True)
bullet(tf, "Навіщо: ознаки RR сильні на ритмічних класах (S), згортка — на "
       "формі (V, F); потоки доповнюють один одного", size=16)


# ============================================ SLIDE 8 — novelty 2
s = prs.slides.add_slide(BLANK)
title_bar(s, "Новизна 2 — бенчмарк багатоформатної стійкості", 8)
add_image_fit(s, os.path.join(RUNF, "format_robustness.png"),
              Inches(7.0), Inches(1.3), Inches(6.0), Inches(5.6))
tf = body_box(s, Inches(0.5), Inches(1.4), Inches(6.4), Inches(5.6))
bullet(tf, "Що робимо:", size=18, bold=True, color=NAVY, first=True)
bullet(tf, "класифікуємо нативний WFDB → конвертуємо у CSV / EDF → знову "
       "класифікуємо → міряємо узгодженість міток і різницю сигналу", size=16, level=1)
bullet(tf, "Навіщо:", size=18, bold=True, color=NAVY)
bullet(tf, "«уніфікована багатоформатна» має сенс лише якщо якість не падає "
       "після конвертації — майже жодна робота цього не перевіряє", size=16, level=1)
bullet(tf, "Результат:", size=18, bold=True, color=ACCENT)
bullet(tf, "CSV — 100 % узгодженості міток; EDF — ≥ 99,3 % (фактично без втрат)",
       size=16, level=1)


# ============================================ SLIDE 9 — results: models
s = prs.slides.add_slide(BLANK)
title_bar(s, "Результати: порівняння моделей (DS2)", 9)
rows = [
    ("Модель", "macro-F1", "Повнота F", "Acc"),
    ("На основі правил", "0,24", "0,00", "0,85"),
    ("Випадковий ліс", "0,37", "0,00", "0,93"),
    ("XGBoost", "0,42", "0,52", "0,87"),
    ("1D-CNN", "0,27", "0,02", "0,67"),
    ("Гібрид (2 відв.)", "0,40", "0,89", "0,77"),
    ("SVM-RBF (2 відв.)", "0,466", "0,90", "0,83"),
]
nrows, ncols = len(rows), 4
gt = s.shapes.add_table(nrows, ncols, Inches(0.7), Inches(1.4),
                        Inches(8.0), Inches(4.5)).table
gt.columns[0].width = Inches(3.5)
for j in range(1, 4):
    gt.columns[j].width = Inches(1.5)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = gt.cell(i, j)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        run = para.add_run(); run.text = val
        best = (i == nrows - 1)  # SVM row
        _set(run, 15, bold=(i == 0 or best),
             color=WHITE if i == 0 else (NAVY if best else DARK))
        if i == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        elif best:
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFA)
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
tf = body_box(s, Inches(9.0), Inches(1.6), Inches(4.1), Inches(5.2))
bullet(tf, "Найкраща модель — SVM-RBF на двох відведеннях:", size=17,
       bold=True, color=NAVY, first=True)
bullet(tf, "macro-F1 = 0,466 — на рівні опублікованих результатів для "
       "міжпацієнтного протоколу", size=15, level=1)
bullet(tf, "Accuracy оманлива: випадковий ліс має 0,93, але майже не виявляє "
       "патологій (macro-F1 лише 0,37)", size=15)


# ============================================ SLIDE 10 — results: rare + format
s = prs.slides.add_slide(BLANK)
title_bar(s, "Результати: рідкісні класи та новизни", 10)
add_image_fit(s, os.path.join(RUNF, "rare_class_recall.png"),
              Inches(7.0), Inches(1.3), Inches(6.0), Inches(5.6))
tf = body_box(s, Inches(0.5), Inches(1.4), Inches(6.4), Inches(5.6))
bullet(tf, "Друге відведення дає найбільший приріст за рідкісними класами",
       size=17, first=True)
bullet(tf, "Гібрид (новизна 1):", size=17, bold=True, color=NAVY)
bullet(tf, "повнота зливних скорочень 0,89 — вище за обидва свої потоки "
       "(0,82 у XGBoost, 0,07 у CNN)", size=15, level=1)
bullet(tf, "пізнє злиття відновлює корисний внесок слабкої мережі; не "
       "перевершує SVM (0,90), але майже наздоганяє його", size=15, level=1)
bullet(tf, "Стійкість форматів (новизна 2):", size=17, bold=True, color=NAVY)
bullet(tf, "CSV — без втрат (100 %); EDF — у межах 16-бітного шуму (≥ 99,3 %)",
       size=15, level=1)


# ============================================ SLIDE 11 — conclusions
s = prs.slides.add_slide(BLANK)
title_bar(s, "Висновки", 11)
cols = [
    ("Досягнуто", NAVY, [
        "ML-конвеєр на MIT-BIH за міжпацієнтним протоколом; 7 моделей",
        "найкраща — SVM-RBF 2 відв.: macro-F1 0,466",
        "гібрид: повнота F 0,89 (вище за обидва потоки)",
        "бенчмарк форматів: CSV 100 %, EDF ≥ 99,3 %",
        "REST-API та CLI",
    ]),
    ("Практична цінність", RGBColor(0x3F, 0x8A, 0x52), [
        "скринінг аритмій, телемедицина",
        "єдиний вхід для різних форматів ЕКГ",
        "повна відтворюваність (усе сценаріями)",
    ]),
    ("Подальший розвиток", RGBColor(0x8A, 0x6A, 0x2A), [
        "навчання на GPU з аугментацією даних",
        "перехід до 12 відведень (база PTB-XL)",
        "розширення бенчмарку на більше форматів",
    ]),
]
x = Inches(0.4)
cw = Inches(4.1)
for name, col, items in cols:
    hd = s.shapes.add_shape(1, x, Inches(1.3), cw, Inches(0.6))
    hd.fill.solid(); hd.fill.fore_color.rgb = col; hd.line.fill.background()
    hp = hd.text_frame.paragraphs[0]; hp.alignment = PP_ALIGN.CENTER
    hr = hp.add_run(); hr.text = name; _set(hr, 17, bold=True, color=WHITE)
    tf = body_box(s, x, Inches(2.05), cw, Inches(5.0))
    first = True
    for it in items:
        bullet(tf, it, size=14, first=first, space=8); first = False
    x = x + cw + Inches(0.15)


# ============================================ SLIDE 12 — thanks
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tf = body_box(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.8))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Дякую за увагу!"
_set(r, 40, bold=True, color=WHITE)
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(16)
r = p.add_run(); r.text = "Готовий відповісти на запитання"
_set(r, 20, italic=True, color=RGBColor(0xF0, 0xC4, 0x19))

prs.save(OUT)
print("Saved:", OUT, "| slides:", len(prs.slides._sldIdLst))
